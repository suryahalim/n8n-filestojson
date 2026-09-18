"""S5 fixture: IT asset register continuing the PO-2026-8412 story."""
import openpyxl
from openpyxl.styles import Font

wb = openpyxl.Workbook()

ws = wb.active
ws.title = "Assets"
rows = [
    ["Asset Tag", "Description", "Serial Number", "Location", "Purchase Ref", "Status"],
    ["AST-0112", "Dell PowerStore 300T appliance", "PST300T-77412", "DC Cileungsi Rack A3", "PO-2026-8412", "Accepted"],
    ["AST-0113", "PowerStore expansion enclosure", "PSTEXP-77413", "DC Cileungsi Rack A3", "PO-2026-8412", "Accepted"],
    ["AST-0114", "10GbE SFP+ transceiver", "SFP10G-77420", "DC Cileungsi Rack A3", "PO-2026-8412", "Installed"],
    ["AST-0115", "10GbE SFP+ transceiver", "SFP10G-77421", "DC Cileungsi Rack A4", "PO-2026-8412", "Installed"],
    ["AST-0116", "Rack cabinet 42U + PDU", "RCK42U-118", "DC Cileungsi Row 2", "PO-2026-8412", "Accepted"],
]
for r in rows:
    ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True)

ws2 = wb.create_sheet("Warranty")
for r in [["Asset Tag", "Warranty End", "SLA", "Vendor"],
          ["AST-0112", "2029-11-20", "4x24 on-site", "PT Bukit Limau Teknologi Informasi"],
          ["AST-0113", "2029-11-20", "4x24 on-site", "PT Bukit Limau Teknologi Informasi"],
          ["AST-0116", "2031-11-20", "NBD", "PT Bukit Limau Teknologi Informasi"]]:
    ws2.append(r)
for c in ws2[1]:
    c.font = Font(bold=True)

ws3 = wb.create_sheet("MaintenanceLog")
for r in [["Date", "Asset Tag", "Activity", "PIC", "Evidence"],
          ["2026-11-15", "AST-0112", "Rack mounting & cabling report", "Infra Team", "RPT-MOUNT-1115.pdf"],
          ["2026-11-20", "AST-0112", "Firmware baseline 4.0.2.1 & license upload", "Vendor", "SNIT-4-0-2-1.txt"],
          ["2026-11-29", "AST-0112", "Replication failover test evidence", "Infra Team", "RPT-FAILOVER-1129.pdf"],
          ["2026-12-12", "AST-0112", "Post-implementation review meeting", "IT Manager", "NOTULENS-PIR.pdf"]]:
    ws3.append(r)
for c in ws3[1]:
    c.font = Font(bold=True)

path = "/home/suryahalim/doc-pipeline/tests/fixtures/IT-Asset-Register-PO-8412.xlsx"
wb.save(path)
print("wrote", path)
